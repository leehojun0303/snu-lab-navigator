from __future__ import annotations

import json
from pathlib import Path
from openpyxl import load_workbook

source = Path("/workspace/scratch/192cf3ff2523/upload/snu_lab_discovery_v30_result.xlsx")
destination = Path("/workspace/sites/snu-lab-navigator/dist/data.js")
workbook = load_workbook(source, read_only=True, data_only=True)
sheet = workbook["faculty_research_units_v30"]
headers = [cell.value for cell in sheet[1]]


def rows_by_header(sheet_name):
    source_sheet = workbook[sheet_name]
    source_headers = [cell.value for cell in source_sheet[1]]
    for source_values in source_sheet.iter_rows(min_row=2, values_only=True):
        yield dict(zip(source_headers, source_values))


# Preserve official display assets that are stored outside the v30 research-unit
# sheet. Membership IDs keep this join stable even when names or affiliations vary.
faculty_records = {row.get("faculty_id"): row for row in rows_by_header("faculty_registry")}
memberships = {}
for row in rows_by_header("faculty_entity_membership"):
    memberships.setdefault(row.get("entity_id"), []).append(row.get("faculty_record_id"))

department_links = {}
for row in rows_by_header("unit_coverage_recomputed"):
    key = (row.get("college") or "", row.get("department") or "")
    url = row.get("official_department_url") or row.get("unit_url_final") or ""
    if url:
        department_links[key] = url


def first_value(rows, *columns):
    for row in rows:
        for column in columns:
            value = row.get(column)
            if value:
                return value
    return ""


records = []
for values in sheet.iter_rows(min_row=2, values_only=True):
    row = dict(zip(headers, values))
    linked_records = [
        faculty_records[record_id]
        for record_id in memberships.get(row.get("faculty_entity_id"), [])
        if record_id in faculty_records
    ]
    same_unit_records = [
        record for record in linked_records
        if (record.get("college") or "") == (row.get("college") or "")
        and (record.get("department") or "") == (row.get("department") or "")
    ] or linked_records
    unit_key = (row.get("college") or "", row.get("department") or "")
    department_url = department_links.get(unit_key) or first_value(
        same_unit_records, "department_source_url", "college_source_url"
    )
    records.append({
        "id": row.get("research_unit_id", ""),
        "name": row.get("faculty_name", ""),
        "title": row.get("display_name", ""),
        "college": row.get("college", ""),
        "department": row.get("department", ""),
        "rank": row.get("published_rank", ""),
        "type": row.get("research_unit_type", ""),
        "naming": row.get("naming_status", ""),
        "labs": row.get("verified_lab_names", ""),
        "fields": row.get("research_fields", ""),
        "keywords": row.get("research_keywords", ""),
        "profile": row.get("official_profile_urls", ""),
        "homepage": row.get("published_homepage_urls", ""),
        "photo": first_value(same_unit_records, "faculty_photo_url"),
        "departmentUrl": department_url,
        "guidance": row.get("public_guidance", ""),
    })
workbook.close()
destination.parent.mkdir(parents=True, exist_ok=True)
destination.write_text("window.RESEARCH_UNITS = " + json.dumps(records, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(len(records))
